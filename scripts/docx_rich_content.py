"""Parse lightweight markdown and insert styled Word content (headings, tables, lists)."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from docx.document import Document
    from docx.text.paragraph import Paragraph

BODY_FONT = "Times New Roman"
BODY_SIZE_PT = 12
BULLET_CHAR = "\u2022"

TABLE_SEP_RE = re.compile(r"^\s*\|?(?:\s*:?-+:?\s*\|)+\s*$")
NUMBERED_RE = re.compile(r"^\d+\.\s+")
BULLET_RE = re.compile(r"^[-*+•\u2022]\s+")
EMBEDDED_BULLET_RE = re.compile(r"\n\s*[-*+•\u2022]\s+\S")
EMBEDDED_NUMBER_RE = re.compile(r"\n\s*\d+\.\s+\S")
IMAGE_MD_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
BOLD_INLINE_RE = re.compile(r"\*\*(.+?)\*\*")
INLINE_ESCAPE_RE = re.compile(r"\\([\\`*_\[\]()#+\-.!{}|])")
H3_NUMBER_RE = re.compile(r"^\d+\.\d+\.\d+(?:\s|\t|$)")
H3_LINE_SPLIT_RE = re.compile(r"^(\d+\.\d+\.\d+)(?:\s|\t)(.+)$", re.DOTALL)
H3_TITLE_SEP_RE = re.compile(r"^(.+?)\s+[—–\-:|]\s+(.+)$", re.DOTALL)
H3_TITLE_MAX_CHARS = 52
H3_TITLE_MAX_WORDS = 7
H1_NUMBER_RE = re.compile(r"^\d+\.0(?:\s|\t|$)")
H2_NUMBER_RE = re.compile(r"^\d+\.\d+(?:\s|\t|$)")
H2_PREFIX_RE = re.compile(r"^(\d+\.\d+)\b")
SHALLOW_AGENT_SUB_RE = re.compile(r"^1\.(\d+)(?:\s|\t)(.*)$", re.DOTALL)
HEADING_MD_RE = re.compile(r"^(#{1,6})\s+(.+)$")
DEFAULT_IMAGE_WIDTH_IN = 6.5
# Common Word default numbering ids when template has numbering.xml but no List Bullet style
BULLET_NUM_ID = 10
NUMBER_NUM_ID = 1


@dataclass
class ContentBlock:
    kind: str  # h1, h2, h3, p, bullet, number, table, image, spacer
    text: str | None = None
    rows: list[list[str]] | None = None
    image_path: str | None = None


@dataclass(frozen=True)
class ManualStyleConfig:
    """Resolved from ``templates/manual-styles.json`` and catalog overrides."""

    heading_color_hex: str | None = "0E2841"
    heading_levels: frozenset[int] = frozenset({1, 2, 3})
    heading_color_only_non_default: bool = True
    blank_line_before_heading1: bool = False
    blank_line_before_heading2: bool = True
    heading1_page_break_before: bool = True
    trim_leading_trailing_blank_lines: bool = True
    image_align: str = "center"
    image_caption_enabled: bool = True
    image_caption_font: str = BODY_FONT
    image_caption_size_pt: float = 10
    image_caption_italic: bool = True
    image_caption_prefix: str = "Figure"
    image_caption_use_seq: bool = True
    image_caption_word_style: str = "Caption"
    image_caption_align: str = "center"
    image_caption_blank_line_after: bool = True


DEFAULT_MANUAL_STYLE_CONFIG = ManualStyleConfig()


def _normalize_line_endings(content: str) -> str:
    return content.replace("\r\n", "\n").replace("\r", "\n")


def preprocess_markdown_content(content: str) -> str:
    """Normalize line endings and split embedded list lines onto their own rows."""
    text = _normalize_line_endings(content)
    text = re.sub(r"(\S)\n(\s*[-*+•\u2022]\s+)", r"\1\n\n\2", text)
    text = re.sub(r"(\S)\n(\s*\d+\.\s+)", r"\1\n\n\2", text)
    return text


def _paragraph_has_embedded_list(text: str) -> bool:
    norm = _normalize_line_endings(text)
    return bool(EMBEDDED_BULLET_RE.search(norm) or EMBEDDED_NUMBER_RE.search(norm))


def _parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return "|" in s and not TABLE_SEP_RE.match(s)


def _numbered_heading_kind(title: str) -> str | None:
    """Map manual-style numbers (2.0 → h1, 3.1 → h2, 3.1.1 → h3)."""
    t = title.strip()
    if not t:
        return None
    if H3_NUMBER_RE.match(t):
        return "h3"
    if H1_NUMBER_RE.match(t):
        return "h1"
    if H2_NUMBER_RE.match(t):
        return "h2"
    return None


def _resolve_heading_kind(markdown_hashes: int, title: str) -> str:
    numbered = _numbered_heading_kind(title)
    if numbered:
        return numbered
    if markdown_hashes == 1:
        return "h1"
    if markdown_hashes == 2:
        return "h2"
    return "h3"


def _h2_prefix_from_title(title: str) -> str | None:
    m = H2_PREFIX_RE.match(title.strip())
    return m.group(1) if m else None


def _is_shallow_agent_subsection(text: str) -> bool:
    """Agent shorthand ``1.1`` / ``1.2`` under a real ``4.3`` subsection."""
    return bool(SHALLOW_AGENT_SUB_RE.match(text.strip()))


def _already_h3_under_prefix(h2_prefix: str, text: str) -> bool:
    t = text.strip()
    return t.startswith(f"{h2_prefix}.") and bool(H3_NUMBER_RE.match(t))


def _remap_shallow_subsection(h2_prefix: str, text: str, *, seq: int) -> str | None:
    """Map agent shorthand ``1.N …`` under ``6.2`` → ``6.2.{seq} …`` (sequential, not ``6.2.18``)."""
    t = text.strip()
    if not t or _already_h3_under_prefix(h2_prefix, t):
        return None
    m = SHALLOW_AGENT_SUB_RE.match(t)
    if not m:
        return None
    rest = m.group(2).strip()
    head = f"{h2_prefix}.{seq}"
    if not rest:
        return head
    sep = "\t" if "\t" in t else " "
    return f"{head}{sep}{rest}"


def _normalize_subsection_numbers(blocks: list[ContentBlock]) -> list[ContentBlock]:
    """Remap agent shorthand ``1.1`` under ``4.3`` to ``4.3.1`` — do **not** promote markdown ``1.`` lists to Heading 3."""
    out: list[ContentBlock] = []
    h2_prefix: str | None = None
    h3_seq = 0
    under_h2 = False

    for block in blocks:
        if block.kind == "h1":
            h2_prefix = None
            h3_seq = 0
            under_h2 = False
            out.append(block)
            continue

        if block.kind == "h2" and block.text:
            remapped = (
                _remap_shallow_subsection(h2_prefix, block.text, seq=h3_seq + 1)
                if h2_prefix and under_h2
                else None
            )
            if remapped:
                h3_seq += 1
                out.append(ContentBlock(kind="h3", text=remapped))
                continue
            prefix = _h2_prefix_from_title(block.text)
            h3_seq = 0
            if prefix and not _is_shallow_agent_subsection(block.text):
                h2_prefix = prefix
                under_h2 = True
            else:
                under_h2 = False
            out.append(block)
            continue

        if block.kind == "h3":
            h3_seq = 0
            under_h2 = False
            out.append(block)
            continue

        if block.text and h2_prefix and under_h2 and block.kind in ("h2", "p"):
            remapped = _remap_shallow_subsection(h2_prefix, block.text, seq=h3_seq + 1)
            if remapped:
                h3_seq += 1
                out.append(ContentBlock(kind="h3", text=remapped))
                continue

        out.append(block)

    return out


def _is_short_h3_title(rest: str) -> bool:
    """Short label stays on one Heading 3 line for TOC (e.g. ``E-Stop Guidelines``)."""
    if not rest:
        return True
    words = rest.split()
    if len(rest) > H3_TITLE_MAX_CHARS or len(words) > H3_TITLE_MAX_WORDS:
        return False
    if rest.endswith(".") and len(words) > 4:
        return False
    return True


def _compact_h3_title_from_long(rest: str) -> str:
    """Fallback index label when a long step was written on one ``###`` line."""
    words = rest.split()
    taken: list[str] = []
    length = 0
    for w in words:
        add = len(w) + (1 if taken else 0)
        if len(taken) >= 4 or (length + add > 42 and taken):
            break
        taken.append(w)
        length += add
    return " ".join(taken) if taken else rest[:42].rsplit(" ", 1)[0]


def _split_h3_number_and_body(text: str) -> tuple[str, str | None]:
    """Heading 3 = ``N.N.N Short title``; long procedure text → Normal body below."""
    t = text.strip()
    m = H3_LINE_SPLIT_RE.match(t)
    if not m:
        return t, None
    number, rest = m.group(1), m.group(2).strip()
    if not rest:
        return number, None

    sep = H3_TITLE_SEP_RE.match(rest)
    if sep:
        title_part, body_part = sep.group(1).strip(), sep.group(2).strip()
        if title_part and body_part:
            return f"{number} {title_part}", body_part

    if _is_short_h3_title(rest):
        return f"{number} {rest}", None

    short = _compact_h3_title_from_long(rest)
    return f"{number} {short}", rest


def _trim_edge_spacers(blocks: list[ContentBlock], *, enabled: bool = True) -> list[ContentBlock]:
    """Drop leading/trailing blank-line spacers (e.g. before ``# 4.0``)."""
    if not enabled:
        return blocks
    while blocks and blocks[0].kind == "spacer":
        blocks.pop(0)
    while blocks and blocks[-1].kind == "spacer":
        blocks.pop()
    return blocks


def _normalize_hex_color(value: str | None) -> str | None:
    if not value:
        return None
    s = str(value).strip().lstrip("#")
    if len(s) == 6 and re.fullmatch(r"[0-9a-fA-F]{6}", s):
        return s.upper()
    return None


def _heading_level_from_style(style_name: str) -> int | None:
    if not style_name.startswith("Heading "):
        return None
    try:
        return int(style_name.rsplit(" ", 1)[-1])
    except ValueError:
        return None


def _run_uses_default_color(run) -> bool:
    from docx.shared import RGBColor

    rgb = run.font.color.rgb
    if rgb is None:
        return True
    return rgb == RGBColor(0, 0, 0)


def _apply_heading_color_to_paragraph(para, hex_color: str, *, only_non_default: bool) -> None:
    from docx.shared import RGBColor

    target = RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )
    for run in para.runs:
        if only_non_default and not _run_uses_default_color(run):
            continue
        run.font.color.rgb = target


def _expand_h3_number_body(blocks: list[ContentBlock]) -> list[ContentBlock]:
    expanded: list[ContentBlock] = []
    for block in blocks:
        if block.kind == "h3" and block.text:
            number, body = _split_h3_number_and_body(block.text)
            expanded.append(ContentBlock(kind="h3", text=number))
            if body:
                expanded.append(ContentBlock(kind="p", text=body))
        else:
            expanded.append(block)
    return expanded


def parse_markdown_blocks(
    content: str,
    *,
    style_config: ManualStyleConfig | None = None,
) -> list[ContentBlock]:
    """Split section text into headings, paragraphs, lists, and pipe tables."""
    lines = preprocess_markdown_content(content).split("\n")
    blocks: list[ContentBlock] = []
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            if blocks and blocks[-1].kind != "spacer":
                blocks.append(ContentBlock(kind="spacer"))
            i += 1
            continue

        if _is_table_row(stripped) and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].strip()):
            rows = [_parse_table_row(stripped)]
            i += 2
            while i < len(lines) and _is_table_row(lines[i]):
                rows.append(_parse_table_row(lines[i]))
                i += 1
            blocks.append(ContentBlock(kind="table", rows=rows))
            continue

        img_match = IMAGE_MD_RE.match(stripped)
        if img_match:
            blocks.append(
                ContentBlock(
                    kind="image",
                    text=img_match.group(1).strip(),
                    image_path=img_match.group(2).strip(),
                )
            )
            i += 1
            continue

        heading_match = HEADING_MD_RE.match(stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            blocks.append(
                ContentBlock(kind=_resolve_heading_kind(level, title), text=title),
            )
        elif (footnote := _parse_escaped_footnote_line(stripped)):
            blocks.append(ContentBlock(kind="p", text=footnote))
        elif BULLET_RE.match(stripped):
            blocks.append(ContentBlock(kind="bullet", text=BULLET_RE.sub("", stripped, count=1).strip()))
        elif NUMBERED_RE.match(stripped):
            blocks.append(ContentBlock(kind="number", text=NUMBERED_RE.sub("", stripped, count=1).strip()))
        else:
            blocks.append(ContentBlock(kind="p", text=stripped))

        i += 1

    cfg = style_config or DEFAULT_MANUAL_STYLE_CONFIG
    blocks = _normalize_subsection_numbers(_expand_multiline_blocks(blocks, style_config=cfg))
    blocks = _expand_h3_number_body(blocks)
    return _trim_edge_spacers(blocks, enabled=cfg.trim_leading_trailing_blank_lines)


def _expand_multiline_blocks(
    blocks: list[ContentBlock],
    *,
    style_config: ManualStyleConfig | None = None,
) -> list[ContentBlock]:
    """Re-parse paragraph blocks that still contain embedded line breaks."""
    expanded: list[ContentBlock] = []
    for block in blocks:
        if block.kind == "p" and block.text and "\n" in block.text:
            expanded.extend(parse_markdown_blocks(block.text, style_config=style_config))
        else:
            expanded.append(block)
    return expanded


def _caption_font(run, cfg: ManualStyleConfig) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    name = cfg.image_caption_font or BODY_FONT
    run.font.name = name
    run.font.size = Pt(cfg.image_caption_size_pt)
    run.font.italic = cfg.image_caption_italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:cs"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def _style_caption_runs(para, cfg: ManualStyleConfig) -> None:
    for run in para.runs:
        _caption_font(run, cfg)


def _append_seq_figure_field(paragraph) -> None:
    """Word SEQ Figure field so Insert → Table of Figures can find captions."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " SEQ Figure \\* ARABIC ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)  # noqa: SLF001


def _resolve_image_caption_text(caption: str | None, image_path: str) -> str:
    text = (caption or "").strip()
    if text:
        return text
    if image_path:
        stem = Path(image_path).stem.replace("_", " ").replace("-", " ")
        if stem.strip():
            return stem.strip()
    return "Image"


def _insert_figure_caption(
    doc: Document,
    insert_after,
    caption_text: str,
    *,
    style_config: ManualStyleConfig,
    image_path: str = "",
) -> Any:
    style = (
        style_config.image_caption_word_style
        if _style_exists(doc, style_config.image_caption_word_style)
        else _style_name(doc, "Normal")
    )
    para = doc.add_paragraph(style=style)
    prefix = (style_config.image_caption_prefix or "Figure").strip()
    body = _resolve_image_caption_text(caption_text, image_path)

    if style_config.image_caption_use_seq:
        lead = para.add_run(f"{prefix} ")
        _caption_font(lead, style_config)
        _append_seq_figure_field(para)
        sep = para.add_run(f": {body}")
        _caption_font(sep, style_config)
    else:
        line = para.add_run(f"{prefix}: {body}")
        _caption_font(line, style_config)

    align = (style_config.image_caption_align or "center").strip().lower()
    if align == "center":
        _set_paragraph_alignment(para, "center")

    return _move_element_after(para._element, insert_after)  # noqa: SLF001


def _body_font(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = BODY_FONT
    run.font.size = Pt(BODY_SIZE_PT)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:cs"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), BODY_FONT)


def _style_body_paragraph(para) -> None:
    if not para.runs:
        if para.text:
            _fill_paragraph_with_inline_markdown(para, para.text, apply_body_font=True)
        return
    for run in para.runs:
        _body_font(run)


def _unescape_inline_markdown(text: str) -> str:
    """Turn ``\\*``, ``\\_``, etc. into literal characters (common in staged footnotes)."""
    return INLINE_ESCAPE_RE.sub(r"\1", text)


def _parse_escaped_footnote_line(stripped: str) -> str | None:
    """Line-start ``\\* `` in staged markdown → Normal paragraph with literal ``* `` prefix."""
    if not stripped.startswith("\\*"):
        return None
    rest = stripped[2:].lstrip()
    return f"* {rest}" if rest else "*"


def _add_inline_markdown_runs(para, text: str, *, apply_body_font: bool = False) -> None:
    if not text:
        return
    text = _unescape_inline_markdown(text)
    if "**" not in text:
        run = para.add_run(text)
        if apply_body_font:
            _body_font(run)
        return
    pos = 0
    for match in BOLD_INLINE_RE.finditer(text):
        if match.start() > pos:
            run = para.add_run(text[pos : match.start()])
            if apply_body_font:
                _body_font(run)
        run = para.add_run(match.group(1))
        run.bold = True
        if apply_body_font:
            _body_font(run)
        pos = match.end()
    if pos < len(text):
        run = para.add_run(text[pos:])
        if apply_body_font:
            _body_font(run)


def _fill_paragraph_with_inline_markdown(para, text: str, *, apply_body_font: bool = False) -> None:
    para.text = ""
    _add_inline_markdown_runs(para, text, apply_body_font=apply_body_font)


def _style_cell(cell) -> None:
    for para in cell.paragraphs:
        _style_body_paragraph(para)


def _style_exists(doc: Document, name: str) -> bool:
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def _style_name(doc: Document, preferred: str, fallback: str = "Normal") -> str:
    return preferred if _style_exists(doc, preferred) else fallback


def _table_style_name(doc: Document) -> str | None:
    for candidate in ("Table Grid", "Normal Table", "Light Grid", "Medium Grid 1"):
        if _style_exists(doc, candidate):
            return candidate
    return None


CENTER_COLUMN_HEADERS = frozenset({"username", "password"})


def _set_table_double_borders(table) -> None:
    """double-rule grid (controls-manual style): double lines on outer border and between all cells."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl  # noqa: SLF001
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "double")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")


def _set_paragraph_alignment(para, alignment: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if alignment == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        val = "center"
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        val = "left"
    p_pr = para._element.get_or_add_pPr()  # noqa: SLF001
    jc = p_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        p_pr.append(jc)
    jc.set(qn("w:val"), val)


def _apply_table_column_alignment(table) -> None:
    if not table.rows:
        return
    headers = [c.text.strip().lower() for c in table.rows[0].cells]
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            header = headers[ci] if ci < len(headers) else ""
            if ri == 0:
                align = "left"
            elif header in CENTER_COLUMN_HEADERS:
                align = "center"
            else:
                align = "left"
            for para in cell.paragraphs:
                _set_paragraph_alignment(para, align)


def _bold_table_row(table, row_index: int = 0) -> None:
    for cell in table.rows[row_index].cells:
        for para in cell.paragraphs:
            if not para.runs:
                if para.text:
                    run = para.add_run(para.text)
                    para.text = ""
                    run.bold = True
                    _body_font(run)
                continue
            for run in para.runs:
                run.bold = True


def _apply_manual_table_format(table, doc: Document) -> None:
    """All injected tables: double borders, bold header row, body font in cells."""
    style_name = _table_style_name(doc)
    if style_name:
        try:
            table.style = style_name
        except KeyError:
            pass
    _set_table_double_borders(table)
    _bold_table_row(table, 0)
    _apply_table_column_alignment(table)


def _move_element_after(element, insert_after) -> Any:
    element.getparent().remove(element)
    insert_after.addnext(element)
    return element


def _apply_numpr(para, num_id: int, ilvl: int = 0) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = para._element.get_or_add_pPr()  # noqa: SLF001
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl_el)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def _insert_paragraph(
    doc: Document,
    insert_after,
    text: str,
    style: str,
    *,
    apply_body_font: bool = False,
    style_config: ManualStyleConfig | None = None,
) -> Any:
    para = doc.add_paragraph(style=style)
    if text:
        _fill_paragraph_with_inline_markdown(para, text, apply_body_font=apply_body_font)
    cfg = style_config or DEFAULT_MANUAL_STYLE_CONFIG
    level = _heading_level_from_style(style)
    if level == 1 and cfg.heading1_page_break_before:
        para.paragraph_format.page_break_before = True
    if (
        level is not None
        and cfg.heading_color_hex
        and level in cfg.heading_levels
    ):
        _apply_heading_color_to_paragraph(
            para,
            cfg.heading_color_hex,
            only_non_default=cfg.heading_color_only_non_default,
        )
    el = _move_element_after(para._element, insert_after)  # noqa: SLF001
    return el


def _insert_spacer_paragraph(doc: Document, insert_after, style: str) -> Any:
    return _insert_paragraph(doc, insert_after, "", style)


def _insert_bullet_paragraph(doc: Document, insert_after, text: str, *, use_list_style: bool) -> Any:
    if use_list_style:
        return _insert_paragraph(doc, insert_after, text, "List Bullet", apply_body_font=True)

    para = doc.add_paragraph(style="Normal")
    _fill_paragraph_with_inline_markdown(para, text, apply_body_font=True)
    _apply_numpr(para, BULLET_NUM_ID)
    return _move_element_after(para._element, insert_after)


def _insert_number_paragraph(
    doc: Document,
    insert_after,
    text: str,
    number: int,
    *,
    use_list_style: bool,
) -> Any:
    if use_list_style:
        return _insert_paragraph(doc, insert_after, text, "List Number", apply_body_font=True)

    para = doc.add_paragraph(style="Normal")
    _fill_paragraph_with_inline_markdown(para, text, apply_body_font=True)
    _apply_numpr(para, NUMBER_NUM_ID)
    return _move_element_after(para._element, insert_after)


def _insert_table(doc: Document, insert_after, rows: list[list[str]]) -> Any:
    if not rows:
        return insert_after

    nrows = len(rows)
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=nrows, cols=ncols)

    for ri, row_data in enumerate(rows):
        for ci in range(ncols):
            val = row_data[ci] if ci < len(row_data) else ""
            cell = table.rows[ri].cells[ci]
            para = cell.paragraphs[0]
            _fill_paragraph_with_inline_markdown(para, val, apply_body_font=True)
            for extra in cell.paragraphs[1:]:
                extra.text = ""

    _apply_manual_table_format(table, doc)
    return _move_element_after(table._tbl, insert_after)


def _insert_image(
    doc: Document,
    insert_after,
    image_path: str,
    *,
    width_inches: float = DEFAULT_IMAGE_WIDTH_IN,
    caption: str | None = None,
    style_config: ManualStyleConfig | None = None,
) -> Any:
    from docx.shared import Inches

    cfg = style_config or DEFAULT_MANUAL_STYLE_CONFIG
    path = Path(image_path)
    if not path.is_file():
        raise ValueError(f"Image not found: {path}")

    img_para = doc.add_paragraph()
    if cfg.image_align == "center":
        _set_paragraph_alignment(img_para, "center")
    img_para.add_run().add_picture(str(path.resolve()), width=Inches(width_inches))
    insert_after = _move_element_after(img_para._element, insert_after)  # noqa: SLF001

    if cfg.image_caption_enabled:
        cap = _resolve_image_caption_text(caption, image_path)
        insert_after = _insert_figure_caption(
            doc,
            insert_after,
            cap,
            style_config=cfg,
            image_path=image_path,
        )
        if cfg.image_caption_blank_line_after:
            insert_after = _insert_spacer_paragraph(
                doc,
                insert_after,
                _style_name(doc, "Normal"),
            )
    return insert_after


def insert_blocks_after_element(
    doc: Document,
    insert_after,
    blocks: list[ContentBlock],
    *,
    width_inches: float = DEFAULT_IMAGE_WIDTH_IN,
    style_config: ManualStyleConfig | None = None,
) -> Any:
    """Insert content blocks after an XML element; returns last inserted element."""
    cfg = style_config or DEFAULT_MANUAL_STYLE_CONFIG
    h1 = _style_name(doc, "Heading 1")
    h2 = _style_name(doc, "Heading 2")
    h3 = _style_name(doc, "Heading 3")
    normal = _style_name(doc, "Normal")
    use_bullet_style = _style_exists(doc, "List Bullet")
    number_counter = 0
    open_h2_subsection = False
    last_inserted_spacer = False
    skip_spacer_before: set[str] = set()
    if not cfg.blank_line_before_heading1:
        skip_spacer_before.add("h1")
    if not cfg.blank_line_before_heading2:
        skip_spacer_before.add("h2")

    def _reset_number_list() -> None:
        nonlocal number_counter
        number_counter = 0

    def _close_h2_subsection() -> None:
        nonlocal insert_after, open_h2_subsection, last_inserted_spacer
        if open_h2_subsection and not last_inserted_spacer:
            insert_after = _insert_spacer_paragraph(doc, insert_after, normal)
            last_inserted_spacer = True
        open_h2_subsection = False

    for idx, block in enumerate(blocks):
        if block.kind == "spacer":
            nxt = blocks[idx + 1] if idx + 1 < len(blocks) else None
            if nxt is not None and nxt.kind in skip_spacer_before:
                continue
        if block.kind == "h1" and block.text:
            _close_h2_subsection()
            insert_after = _insert_paragraph(
                doc, insert_after, block.text, h1, style_config=cfg,
            )
            open_h2_subsection = False
            last_inserted_spacer = False
            _reset_number_list()
        elif block.kind == "h2" and block.text:
            _close_h2_subsection()
            if cfg.blank_line_before_heading2 and not last_inserted_spacer:
                insert_after = _insert_spacer_paragraph(doc, insert_after, normal)
                last_inserted_spacer = True
            insert_after = _insert_paragraph(
                doc, insert_after, block.text, h2, style_config=cfg,
            )
            open_h2_subsection = True
            last_inserted_spacer = False
            _reset_number_list()
        elif block.kind == "spacer":
            if not last_inserted_spacer:
                insert_after = _insert_spacer_paragraph(doc, insert_after, normal)
                last_inserted_spacer = True
        elif block.kind == "h3" and block.text:
            insert_after = _insert_paragraph(
                doc, insert_after, block.text, h3, style_config=cfg,
            )
            last_inserted_spacer = False
            _reset_number_list()
        elif block.kind == "bullet" and block.text:
            insert_after = _insert_bullet_paragraph(
                doc,
                insert_after,
                block.text,
                use_list_style=use_bullet_style,
            )
            last_inserted_spacer = False
            _reset_number_list()
        elif block.kind == "number" and block.text:
            number_counter += 1
            step_text = NUMBERED_RE.sub("", block.text.strip(), count=1).strip() or block.text.strip()
            insert_after = _insert_paragraph(
                doc,
                insert_after,
                f"{number_counter}. {step_text}",
                normal,
                apply_body_font=True,
                style_config=cfg,
            )
            last_inserted_spacer = False
        elif block.kind == "p" and block.text:
            insert_after = _insert_paragraph(
                doc, insert_after, block.text, normal, apply_body_font=True, style_config=cfg,
            )
            last_inserted_spacer = False
            _reset_number_list()
        elif block.kind == "table" and block.rows:
            insert_after = _insert_table(doc, insert_after, block.rows)
            last_inserted_spacer = False
            _reset_number_list()
        elif block.kind == "image" and block.image_path:
            insert_after = _insert_image(
                doc,
                insert_after,
                block.image_path,
                width_inches=width_inches,
                caption=block.text,
                style_config=cfg,
            )
            last_inserted_spacer = False
            _reset_number_list()

    _close_h2_subsection()
    return insert_after


def insert_rich_content_after(anchor_para: Paragraph, content: str) -> dict[str, int]:
    """Insert parsed blocks immediately after anchor paragraph."""
    doc = anchor_para.part.document
    blocks = parse_markdown_blocks(content)
    counts: dict[str, int] = {
        "h1": 0,
        "h2": 0,
        "h3": 0,
        "p": 0,
        "bullet": 0,
        "number": 0,
        "table": 0,
        "image": 0,
        "spacer": 0,
    }
    for block in blocks:
        if block.kind in counts:
            counts[block.kind] += 1
    insert_blocks_after_element(doc, anchor_para._element, blocks)  # noqa: SLF001
    return counts
