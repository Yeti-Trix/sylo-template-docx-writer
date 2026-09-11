"""Shared docx helpers for sylo-template-docx-writer."""

from __future__ import annotations

import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx_rich_content import (
    ManualStyleConfig,
    _fill_paragraph_with_inline_markdown,
    _insert_image,
    _insert_paragraph,
    _normalize_hex_color,
    _paragraph_has_embedded_list,
    _resolve_image_caption_text,
    _style_body_paragraph,
    _style_cell,
    _style_name,
    insert_blocks_after_element,
    parse_markdown_blocks,
)

PLACEHOLDER_RE = re.compile(r"\{([A-Z0-9_]+)\}")
SECTION_H1_RE = re.compile(r"^\d+\.0\b")

MARKER_PREFIX = "SYLO-SECTION"


def open_document(path: Path):
    from docx import Document

    return Document(str(path))


def _marker_text(placeholder: str, which: str) -> str:
    return f"[[{MARKER_PREFIX}:{placeholder.upper()}:{which}]]"


def _set_vanish(rpr) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if rpr.find(qn("w:vanish")) is None:
        rpr.append(OxmlElement("w:vanish"))


def _make_marker_paragraph(doc, placeholder: str, which: str):
    """Hidden (vanish) paragraph that durably tags the start/end of a section."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    para = doc.add_paragraph()
    run = para.add_run(_marker_text(placeholder, which))
    _set_vanish(run._element.get_or_add_rPr())  # noqa: SLF001
    p_pr = para._p.get_or_add_pPr()  # noqa: SLF001
    mark_rpr = p_pr.find(qn("w:rPr"))
    if mark_rpr is None:
        mark_rpr = OxmlElement("w:rPr")
        p_pr.append(mark_rpr)
    _set_vanish(mark_rpr)
    return para


def _find_marker_element(doc, placeholder: str, which: str):
    target = _marker_text(placeholder, which)
    for para in doc.paragraphs:
        if para.text.strip() == target:
            return para._element  # noqa: SLF001
    return None


def _move_after(element, anchor):
    element.getparent().remove(element)
    anchor.addnext(element)
    return element


def _is_deletable_empty_para(para) -> bool:
    if MARKER_PREFIX in para.text:
        return False
    return not para.text.strip()


def _drop_empty_paragraph_before_element(doc, element) -> None:
    """Remove a blank Normal paragraph left from a consumed ``{SNN_SECTION}`` line."""
    prev = element.getprevious()
    if prev is None:
        return
    for para in doc.paragraphs:
        if para._element is prev and _is_deletable_empty_para(para):
            _delete_paragraph(para)
            break


def _place_section_begin_marker(doc, ph: str, para, tag: str):
    """Insert hidden BEGIN marker; drop empty placeholder-only paragraph when possible."""
    before, _, after = para.text.partition(tag)
    para.text = (before + after).strip()
    el = para._element  # noqa: SLF001
    if para.text.strip():
        begin_para = _make_marker_paragraph(doc, ph, "BEGIN")
        begin_el = begin_para._element  # noqa: SLF001
        _move_after(begin_el, el)
        return begin_el

    parent = el.getparent()
    prev = el.getprevious()
    nxt = el.getnext()
    parent.remove(el)
    begin_para = _make_marker_paragraph(doc, ph, "BEGIN")
    begin_el = begin_para._element  # noqa: SLF001
    if prev is not None:
        _move_after(begin_el, prev)
        return begin_el
    if nxt is not None:
        begin_el.getparent().remove(begin_el)
        parent.insert(parent.index(nxt), begin_el)
        return begin_el
    return begin_el


def _clear_between(begin_el, end_el) -> int:
    removed = 0
    node = begin_el.getnext()
    while node is not None and node is not end_el:
        nxt = node.getnext()
        node.getparent().remove(node)
        removed += 1
        node = nxt
    return removed


def _normalize_section_slug(placeholder: str) -> str:
    ph = placeholder.strip().upper()
    if ph.startswith("{") and ph.endswith("}"):
        ph = ph[1:-1]
    if not ph.endswith("_SECTION"):
        raise ValueError(f"Section slug must end with _SECTION, e.g. S12_SECTION (got {placeholder!r})")
    return ph


def _heading_number_for_slug(ph: str) -> str | None:
    match = re.match(r"S(\d+)_SECTION", ph)
    if not match:
        return None
    return f"{int(match.group(1))}.0"


def _anchor_after_section(doc, after_placeholder: str):
    """Return the XML element after which a new top-level section may be inserted."""
    ph = _normalize_section_slug(after_placeholder)
    end_el = _find_marker_element(doc, ph, "END")
    if end_el is not None:
        return end_el

    tag = f"{{{ph}}}"
    section_prefix = _heading_number_for_slug(ph)
    collecting = False
    anchor = None
    for para in doc.paragraphs:
        el = para._element  # noqa: SLF001
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if tag in para.text or text == _marker_text(ph, "BEGIN"):
            collecting = True
            anchor = el
            continue
        if not collecting and section_prefix and style == "Heading 1" and text.startswith(section_prefix):
            collecting = True
            anchor = el
            continue
        if collecting:
            if style == "Heading 1" and SECTION_H1_RE.match(text):
                if section_prefix and text.startswith(section_prefix):
                    anchor = el
                    continue
                if tag not in para.text:
                    break
            anchor = el
    if anchor is None:
        raise ValueError(
            f"Could not locate section {tag} in draft — check after_placeholder or Heading 1 ({section_prefix})."
        )
    return anchor


def add_section_to_draft(
    draft_path: Path,
    state_path: Path,
    *,
    section_number: str,
    title: str,
    placeholder: str,
    after_placeholder: str,
) -> dict[str, Any]:
    """Insert a new ``{SNN_SECTION}`` placeholder into draft.docx (Heading 1 comes from staging)."""
    ph = _normalize_section_slug(placeholder)
    tag = f"{{{ph}}}"
    number = section_number.strip()
    section_title = title.strip()
    if not number or not section_title:
        raise ValueError("section_number and title required")

    doc = open_document(draft_path)
    if tag in discover_placeholders(draft_path):
        raise ValueError(f"Section already exists in draft: {tag}")

    anchor = _anchor_after_section(doc, after_placeholder)
    normal = _style_name(doc, "Normal")
    _insert_paragraph(doc, anchor, tag, normal)
    doc.save(str(draft_path))

    state = load_state(state_path)
    placeholders = state.setdefault("placeholders", [])
    pending = state.setdefault("pending", [])
    if ph not in placeholders:
        placeholders.append(ph)
        placeholders.sort()
    if ph not in pending and ph not in state.get("completed", []):
        pending.append(ph)
        pending.sort()

    state_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
    return {
        "placeholder": ph,
        "stage_first_line": f"# {number} {section_title}",
        "after": _normalize_section_slug(after_placeholder),
        "draft_path": str(draft_path.resolve()),
        "state_path": str(state_path.resolve()),
    }


def fill_cover_placeholders(
    path: Path,
    *,
    project_title: str | None = None,
    job_number: str | None = None,
    manual_type: str | None = None,
) -> dict[str, Any]:
    """Replace ``{S00_COVER_*}`` tags in draft/template while keeping Word run formatting."""
    replacements: list[tuple[str, str]] = []
    if project_title is not None:
        replacements.append(("{S00_COVER_PROJECT_TITLE}", project_title.strip()))
    if job_number is not None:
        replacements.append(("{S00_COVER_JOB_NUMBER}", job_number.strip()))
    if manual_type is not None:
        replacements.append(("{S00_COVER_MANUAL_TYPE}", manual_type.strip()))
    if not replacements:
        raise ValueError("At least one of project_title, job_number, manual_type is required.")

    doc = open_document(path)
    filled: list[str] = []
    for para in doc.paragraphs:
        for tag, value in replacements:
            if tag in para.text and replace_substring_preserving_runs(para, tag, value):
                key = tag.strip("{}")
                if key not in filled:
                    filled.append(key)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for tag, value in replacements:
                        if tag in para.text and replace_substring_preserving_runs(para, tag, value):
                            key = tag.strip("{}")
                            if key not in filled:
                                filled.append(key)

    if not filled:
        raise ValueError("No cover placeholders found in document.")

    doc.save(str(path))
    return {"filled": filled, "path": str(path.resolve())}


def discover_placeholders(path: Path) -> list[str]:
    doc = open_document(path)
    found: set[str] = set()
    for para in doc.paragraphs:
        found.update(PLACEHOLDER_RE.findall(para.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    found.update(PLACEHOLDER_RE.findall(para.text))
    return sorted(found)


def replace_substring_preserving_runs(para, old: str, new: str) -> bool:
    """Replace text without resetting run font/size (cover fields, inline tags)."""
    if not old or old not in (para.text or ""):
        return False
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    runs = list(para.runs)
    if not runs:
        return False
    chunks: list[tuple[Any, str]] = [(r, r.text) for r in runs]
    full = "".join(t for _, t in chunks)
    start = full.find(old)
    if start < 0:
        return False
    end = start + len(old)
    pos = 0
    start_run = end_run = 0
    start_off = end_off = 0
    for i, (_, t) in enumerate(chunks):
        nxt = pos + len(t)
        if start < nxt and start >= pos:
            start_run, start_off = i, start - pos
        if end <= nxt and end > pos:
            end_run, end_off = i, end - pos
            break
        pos = nxt
    if start_run == end_run:
        r, t = chunks[start_run]
        r.text = t[:start_off] + new + t[end_off:]
        return True
    left = chunks[start_run][1][:start_off] + new
    chunks[start_run][0].text = left
    for j in range(start_run + 1, end_run):
        chunks[j][0].text = ""
    if end_off < len(chunks[end_run][1]):
        chunks[end_run][0].text = chunks[end_run][1][end_off:]
    else:
        chunks[end_run][0].text = ""
    return True


def _replace_in_paragraph(para, tag: str, replacement: str) -> bool:
    return replace_substring_preserving_runs(para, tag, replacement)


def _workspace_root_from_draft(draft_path: Path) -> Path | None:
    """``…/projects/<id>/draft.docx`` → workspace root (parent of ``projects``)."""
    parent = draft_path.resolve().parent
    if parent.parent.name == "projects":
        return parent.parent.parent
    return None


def _deep_merge_dict(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    out = dict(base)
    for key, val in override.items():
        if isinstance(val, dict) and isinstance(out.get(key), dict):
            out[key] = _deep_merge_dict(out[key], val)
        else:
            out[key] = val
    return out


def _parse_catalog_document_styles(catalog_text: str) -> dict[str, str]:
    """Read key/value rows under ``## Document styles`` in section-catalog.md."""
    match = re.search(r"^##\s+Document styles\s*$", catalog_text, re.MULTILINE | re.IGNORECASE)
    if not match:
        return {}
    rest = catalog_text[match.end() :]
    next_heading = re.search(r"^##\s+\S", rest, re.MULTILINE)
    section = rest[: next_heading.start()] if next_heading else rest
    out: dict[str, str] = {}
    for line in section.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("<!--"):
            continue
        if "|" in stripped and not re.match(r"^\|?\s*[-:]+", stripped):
            cells = [c.strip().strip("`") for c in stripped.split("|") if c.strip()]
            if len(cells) >= 2 and cells[0].lower() not in ("key", "setting"):
                out[cells[0].lower()] = cells[1]
            continue
        kv = re.match(r"^[-*]?\s*`?([a-zA-Z0-9_]+)`?\s*:\s*`?([^`]+)`?\s*$", stripped)
        if kv:
            out[kv.group(1).lower()] = kv.group(2).strip()
    return out


def _catalog_overrides_to_styles(raw: dict[str, str]) -> dict[str, Any]:
    """Map catalog table keys into manual-styles.json shape."""
    out: dict[str, Any] = {}
    color = raw.get("heading_color") or raw.get("heading_colour")
    if color:
        out.setdefault("headings", {})["color"] = color
    levels = raw.get("heading_levels")
    if levels:
        parsed = [int(x) for x in re.findall(r"\d+", levels)]
        if parsed:
            out.setdefault("headings", {})["levels"] = parsed
    apply_nd = raw.get("heading_apply_when_not_default")
    if apply_nd is not None:
        out.setdefault("headings", {})["apply_when_not_default"] = apply_nd.lower() in (
            "1",
            "true",
            "yes",
            "on",
        )
    for key, layout_key in (
        ("blank_line_before_heading1", "blank_line_before_heading1"),
        ("blank_line_before_heading2", "blank_line_before_heading2"),
        ("trim_leading_trailing_blank_lines", "trim_leading_trailing_blank_lines"),
        ("heading1_page_break_before", "heading1_page_break_before"),
    ):
        val = raw.get(key)
        if val is None:
            continue
        out.setdefault("layout", {})[layout_key] = val.lower() in ("1", "true", "yes", "on")
    cap = out.setdefault("image_caption", {})
    if (v := raw.get("image_caption_enabled")) is not None:
        cap["enabled"] = v.lower() in ("1", "true", "yes", "on")
    if (v := raw.get("image_caption_font")) is not None:
        cap["font_name"] = v
    if (v := raw.get("image_caption_size_pt")) is not None:
        try:
            cap["font_size_pt"] = float(v)
        except ValueError:
            pass
    if (v := raw.get("image_caption_italic")) is not None:
        cap["italic"] = v.lower() in ("1", "true", "yes", "on")
    if (v := raw.get("image_caption_prefix")) is not None:
        cap["prefix"] = v
    if (v := raw.get("image_caption_use_seq")) is not None:
        cap["use_seq_field"] = v.lower() in ("1", "true", "yes", "on")
    if (v := raw.get("image_caption_word_style")) is not None:
        cap["word_style"] = v
    return out


def package_manual_styles_path() -> Path:
    return PACKAGE_SKILL_DIR / "manual-styles.json"


def load_package_default_manual_styles() -> dict[str, Any]:
    """Built-in default styles shipped with sylo-template-docx-writer. Bundle projects override these via the bundle's manual-styles.json; legacy bare-.docx projects override via workspace templates/manual-styles.json."""
    path = package_manual_styles_path()
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return {k: v for k, v in data.items() if not str(k).startswith("_")}


def _bundle_path_from_draft(draft_path: Path | None) -> Path | None:
    """Read ``template_bundle_path`` from the project state.json next to the draft."""
    if draft_path is None:
        return None
    state_path = draft_path.resolve().parent / "state.json"
    if not state_path.is_file():
        return None
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    bp = state.get("template_bundle_path")
    if not bp:
        return None
    p = Path(bp)
    return p if p.is_dir() else None


def load_manual_styles_raw(
    *,
    draft_path: Path | None = None,
    workspace_root: Path | None = None,
) -> dict[str, Any]:
    """Package default → template bundle (single source of truth) → legacy workspace ``templates/`` → catalog overrides.

    A project started from a template bundle records ``template_bundle_path`` in
    state.json; that bundle's ``manual-styles.json`` and ``section-catalog.md``
    Document styles are the source of truth. Projects started from a bare .docx
    (no bundle) fall back to the legacy workspace ``templates/`` folder, preserving
    prior behavior exactly.
    """
    merged: dict[str, Any] = load_package_default_manual_styles()

    style_bases: list[Path] = []
    bundle = _bundle_path_from_draft(draft_path)
    if bundle is not None:
        style_bases.append(bundle)
    else:
        root = workspace_root
        if root is None and draft_path is not None:
            root = _workspace_root_from_draft(draft_path)
        if root is not None:
            style_bases.append(root / "templates")

    for base in style_bases:
        ws_path = (base / "manual-styles.json").resolve()
        if ws_path.is_file():
            merged = _deep_merge_dict(merged, json.loads(ws_path.read_text(encoding="utf-8")))
        catalog_path = base / "section-catalog.md"
        if catalog_path.is_file():
            catalog_raw = _parse_catalog_document_styles(catalog_path.read_text(encoding="utf-8"))
            if catalog_raw:
                merged = _deep_merge_dict(merged, _catalog_overrides_to_styles(catalog_raw))
    return merged


def manual_style_config_from_dict(data: dict[str, Any] | None) -> ManualStyleConfig:
    headings = (data or {}).get("headings") or {}
    layout = (data or {}).get("layout") or {}
    caption = (data or {}).get("image_caption") or {}
    image = (data or {}).get("image") or {}
    levels_raw = headings.get("levels", [1, 2, 3])
    levels = frozenset(int(x) for x in levels_raw if isinstance(x, (int, float)) or str(x).isdigit())
    if not levels:
        levels = frozenset({1, 2, 3})
    color = _normalize_hex_color(headings.get("color"))
    cap_size = caption.get("font_size_pt", 10)
    try:
        cap_size_f = float(cap_size)
    except (TypeError, ValueError):
        cap_size_f = 10.0
    return ManualStyleConfig(
        heading_color_hex=color,
        heading_levels=levels,
        heading_color_only_non_default=bool(headings.get("apply_when_not_default", True)),
        blank_line_before_heading1=bool(layout.get("blank_line_before_heading1", False)),
        blank_line_before_heading2=bool(layout.get("blank_line_before_heading2", True)),
        heading1_page_break_before=bool(layout.get("heading1_page_break_before", True)),
        trim_leading_trailing_blank_lines=bool(layout.get("trim_leading_trailing_blank_lines", True)),
        image_align=str(image.get("align", "center")),
        image_caption_enabled=bool(caption.get("enabled", True)),
        image_caption_font=str(caption.get("font_name", "Times New Roman")),
        image_caption_size_pt=cap_size_f,
        image_caption_italic=bool(caption.get("italic", True)),
        image_caption_prefix=str(caption.get("prefix", "Figure")),
        image_caption_use_seq=bool(caption.get("use_seq_field", True)),
        image_caption_word_style=str(caption.get("word_style", "Caption")),
        image_caption_align=str(caption.get("align", "center")),
        image_caption_blank_line_after=bool(caption.get("blank_line_after", True)),
    )


def load_manual_style_config(
    *,
    draft_path: Path | None = None,
    workspace_root: Path | None = None,
) -> ManualStyleConfig:
    return manual_style_config_from_dict(load_manual_styles_raw(draft_path=draft_path, workspace_root=workspace_root))


def inject_section(path: Path, placeholder: str, content: str) -> dict[str, Any]:
    """Inject/replace a section's content (Heading 1/2/3, tables, Times New Roman 12 body).

    Idempotent: the first inject consumes ``{PLACEHOLDER}`` and wraps the content in
    hidden BEGIN/END markers. Subsequent calls find those markers, clear the old
    content, and re-insert — so the agent can edit a section without re-starting the
    project.
    """
    ph = placeholder.upper()
    tag = f"{{{ph}}}"
    doc = open_document(path)
    style_config = load_manual_style_config(draft_path=path)

    begin_el = _find_marker_element(doc, ph, "BEGIN")
    end_el = _find_marker_element(doc, ph, "END")
    if begin_el is not None and end_el is not None:
        _clear_between(begin_el, end_el)
        _drop_empty_paragraph_before_element(doc, begin_el)
        blocks = parse_markdown_blocks(content, style_config=style_config)
        insert_blocks_after_element(doc, begin_el, blocks, style_config=style_config)
        doc.save(str(path))
        return {
            "placeholder": ph,
            "chars": len(content),
            "mode": "revised",
            "path": str(path.resolve()),
            "heading_color": style_config.heading_color_hex,
        }

    for para in list(doc.paragraphs):
        if tag not in para.text:
            continue
        begin = _place_section_begin_marker(doc, ph, para, tag)
        blocks = parse_markdown_blocks(content, style_config=style_config)
        last = insert_blocks_after_element(doc, begin, blocks, style_config=style_config)
        _move_after(_make_marker_paragraph(doc, ph, "END")._element, last)  # noqa: SLF001
        doc.save(str(path))
        return {
            "placeholder": ph,
            "chars": len(content),
            "mode": "injected",
            "path": str(path.resolve()),
            "heading_color": style_config.heading_color_hex,
        }

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if _replace_in_paragraph(para, tag, content.replace("\n", " ")):
                        doc.save(str(path))
                        return {
                            "placeholder": ph,
                            "chars": len(content),
                            "mode": "table-cell",
                            "path": str(path.resolve()),
                        }

    raise ValueError(
        f"Placeholder not found and no existing markers for {tag}. "
        "Section may not exist in this draft — run manual_start_project or check the slug."
    )


def insert_image(
    path: Path,
    placeholder: str,
    image_path: Path,
    *,
    width_inches: float = 6.5,
    caption: str | None = None,
) -> dict[str, Any]:
    tag = f"{{{placeholder.upper()}}}"
    if not image_path.is_file():
        raise ValueError(f"Image not found: {image_path}")

    doc = open_document(path)
    style_config = load_manual_style_config(draft_path=path)
    cap_text = _resolve_image_caption_text(caption, str(image_path))
    replaced = False

    for para in list(doc.paragraphs):
        if tag not in para.text:
            continue
        para.text = para.text.replace(tag, "").strip()
        anchor = para._element  # noqa: SLF001
        if not para.text:
            parent = anchor.getparent()
            prev = anchor.getprevious()
            parent.remove(anchor)
            anchor = prev if prev is not None else anchor
        _insert_image(
            doc,
            anchor,
            str(image_path),
            width_inches=width_inches,
            caption=cap_text,
            style_config=style_config,
        )
        replaced = True
        break

    if not replaced:
        raise ValueError(f"Placeholder not found: {tag}")

    doc.save(str(path))
    return {
        "placeholder": placeholder,
        "image_path": str(image_path.resolve()),
        "width_inches": width_inches,
        "caption": cap_text,
        "path": str(path.resolve()),
    }


HMI_DEFAULT_PLACEHOLDER = "S06_SECTION"


def stage_hmi_screen(
    state_path: Path,
    *,
    title: str,
    body: str,
    image_path: str | None = None,
    placeholder: str = HMI_DEFAULT_PLACEHOLDER,
) -> dict[str, Any]:
    """Append one operator HMI screen to state.json staging (persist across turns)."""
    state = load_state(state_path)
    staging = state.setdefault("staging", {})
    block = staging.setdefault(
        placeholder.upper(),
        {"type": "hmi", "screens": [], "complete": False},
    )
    if block.get("body") and not block.get("screens"):
        raise ValueError(
            f"{placeholder.upper()} already staged as prose (no-HMI §6) — use stage-section only"
        )
    entry = {
        "n": len(block["screens"]) + 1,
        "title": title.strip(),
        "body": body.strip(),
        "image_path": image_path,
        "staged_at": datetime.now(timezone.utc).isoformat(),
    }
    block["screens"].append(entry)
    state_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
    return {
        "placeholder": placeholder.upper(),
        "screen_count": len(block["screens"]),
        "staged": entry,
    }


def get_hmi_staging(state_path: Path, placeholder: str = HMI_DEFAULT_PLACEHOLDER) -> dict[str, Any]:
    state = load_state(state_path)
    block = state.get("staging", {}).get(placeholder.upper(), {"screens": [], "complete": False})
    return {
        "placeholder": placeholder.upper(),
        "complete": block.get("complete", False),
        "screen_count": len(block.get("screens", [])),
        "screens": block.get("screens", []),
    }


def finalize_hmi_section(
    draft_path: Path,
    placeholder: str,
    screens: list[dict[str, Any]],
    *,
    width_inches: float = 6.5,
) -> dict[str, Any]:
    """Replace {PLACEHOLDER} with numbered subsections (title, body, image each)."""
    ph = placeholder.upper()
    tag = f"{{{ph}}}"
    doc = open_document(draft_path)

    begin_el = _find_marker_element(doc, ph, "BEGIN")
    end_el = _find_marker_element(doc, ph, "END")
    if begin_el is not None and end_el is not None:
        _clear_between(begin_el, end_el)
        insert_after = begin_el
    else:
        anchor_para = None
        for para in doc.paragraphs:
            if tag in para.text:
                anchor_para = para
                para.text = para.text.replace(tag, "")
                break
        if anchor_para is None:
            raise ValueError(f"Placeholder not found: {tag}")
        anchor = anchor_para._element  # noqa: SLF001
        begin_el = _move_after(_make_marker_paragraph(doc, ph, "BEGIN")._element, anchor)  # noqa: SLF001
        insert_after = begin_el

    h3_style = _style_name(doc, "Heading 3")
    style_config = load_manual_style_config(draft_path=draft_path)

    for screen in screens:
        n = screen.get("n") or 1
        title = str(screen.get("title", "")).strip()
        body = str(screen.get("body", "")).strip()
        image_path = screen.get("image_path")

        insert_after = _insert_paragraph(
            doc,
            insert_after,
            f"6.{n}\t{title}",
            h3_style,
            style_config=style_config,
        )

        if image_path:
            ip = Path(str(image_path))
            if not ip.is_file():
                raise ValueError(f"Image not found: {ip}")
            cap = title or f"Operator screen 6.{n}"
            insert_after = _insert_image(
                doc,
                insert_after,
                str(ip),
                width_inches=width_inches,
                caption=cap,
                style_config=style_config,
            )

        if body:
            blocks = parse_markdown_blocks(body, style_config=style_config)
            insert_after = insert_blocks_after_element(
                doc,
                insert_after,
                blocks,
                width_inches=width_inches,
                style_config=style_config,
            )

    _move_after(_make_marker_paragraph(doc, ph, "END")._element, insert_after)  # noqa: SLF001

    doc.save(str(draft_path))
    return {
        "placeholder": ph,
        "screens_injected": len(screens),
        "path": str(draft_path.resolve()),
    }


def finalize_hmi_from_state(
    state_path: Path,
    *,
    draft_path: Path | None = None,
    placeholder: str = HMI_DEFAULT_PLACEHOLDER,
    width_inches: float = 6.5,
) -> dict[str, Any]:
    state = load_state(state_path)
    draft = Path(draft_path or state["draft_path"])
    ph = placeholder.upper()
    block = state.get("staging", {}).get(ph)
    if not block or not block.get("screens"):
        raise ValueError(f"No staged screens for {ph}")

    result = finalize_hmi_section(
        draft,
        ph,
        block["screens"],
        width_inches=width_inches,
    )
    block["complete"] = True
    state_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
    mark_placeholder_done(state_path, ph)
    result["state_updated"] = True
    return result


def stage_section(state_path: Path, placeholder: str, body: str) -> dict[str, Any]:
    """Store one approved prose section in state.json (does not touch draft.docx)."""
    ph = placeholder.upper()
    state = load_state(state_path)
    staging = state.setdefault("staging", {})
    existing = staging.get(ph)
    if existing and existing.get("screens"):
        raise ValueError(
            f"{ph} has staged HMI screens — finish with manual_finalize_hmi / manual_build_draft, "
            "not stage-section"
        )

    staging[ph] = {
        "type": "prose",
        "body": body.strip(),
        "staged_at": datetime.now(timezone.utc).isoformat(),
    }
    state_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
    return {"placeholder": ph, "staged": True, "chars": len(body.strip())}


def revise_section(
    state_path: Path,
    placeholder: str,
    body: str,
    *,
    draft_path: Path | None = None,
) -> dict[str, Any]:
    """Re-stage one section and re-inject it into draft.docx in place.

    Requires the section to have been built once already (BEGIN/END markers present)
    or its ``{PLACEHOLDER}`` to still exist. Lets the agent edit a section's text
    without re-running manual_start_project.
    """
    ph = placeholder.upper()
    stage_section(state_path, ph, body)
    state = load_state(state_path)
    draft = Path(draft_path or state["draft_path"])
    result = inject_section(draft, ph, body)
    mark_placeholder_done(state_path, ph)
    result["state_updated"] = True
    return result


def list_staging(state_path: Path) -> dict[str, Any]:
    state = load_state(state_path)
    staging = state.get("staging", {})
    prose: list[str] = []
    hmi: dict[str, Any] | None = None

    for ph, block in staging.items():
        if block.get("screens") is not None:
            hmi = {
                "placeholder": ph,
                "screen_count": len(block.get("screens", [])),
                "complete": block.get("complete", False),
            }
        elif block.get("body"):
            prose.append(ph)

    pending = state.get("pending", [])
    unstaged = [p for p in pending if p not in prose and (not hmi or p != hmi.get("placeholder"))]
    return {
        "prose_staged": sorted(prose),
        "hmi": hmi,
        "pending_not_staged": unstaged,
        "access_table": state.get("access_table", "pending"),
    }


def repair_embedded_lists(path: Path) -> dict[str, Any]:
    """Split legacy single-paragraph sections that contain ``\\n- item`` into Word list paragraphs."""
    doc = open_document(path)
    skip_styles = {"Heading 1", "Heading 2", "Heading 3", "Title", "toc 1", "toc 2", "toc 3"}
    repaired = 0
    list_blocks_inserted = 0

    for para in list(doc.paragraphs):
        if para.style and para.style.name in skip_styles:
            continue
        if MARKER_PREFIX in para.text:
            continue
        text = para.text
        if not text or not _paragraph_has_embedded_list(text):
            continue

        style_config = load_manual_style_config(draft_path=path)
        blocks = parse_markdown_blocks(text, style_config=style_config)
        list_blocks = [b for b in blocks if b.kind in ("bullet", "number")]
        if not list_blocks:
            continue

        lead = next((b for b in blocks if b.kind == "p"), None)
        rest = [b for b in blocks if b.kind != "p" or b is not lead]

        if lead and lead.text:
            _fill_paragraph_with_inline_markdown(para, lead.text, apply_body_font=True)
        else:
            para.text = ""

        if rest:
            insert_blocks_after_element(doc, para._element, rest, style_config=style_config)  # noqa: SLF001
            list_blocks_inserted += len(rest)

        repaired += 1

    doc.save(str(path))
    return {
        "paragraphs_repaired": repaired,
        "list_blocks_inserted": list_blocks_inserted,
        "path": str(path.resolve()),
    }


def build_draft_from_staging(
    state_path: Path,
    *,
    draft_path: Path | None = None,
    include_hmi: bool = True,
) -> dict[str, Any]:
    """Inject all staged sections into draft.docx (prose slugs, then HMI if staged)."""
    state = load_state(state_path)
    draft = Path(draft_path or state["draft_path"])
    injected: list[str] = []

    for ph in state.get("placeholders", []):
        block = state.get("staging", {}).get(ph)
        if not block or block.get("screens") is not None:
            continue
        body = block.get("body")
        if not body:
            continue
        inject_section(draft, ph, body)
        mark_placeholder_done(state_path, ph)
        injected.append(ph)

    if include_hmi:
        hmi_block = state.get("staging", {}).get(HMI_DEFAULT_PLACEHOLDER)
        if hmi_block and hmi_block.get("screens") and not hmi_block.get("complete"):
            finalize_hmi_from_state(state_path, draft_path=draft)
            injected.append(HMI_DEFAULT_PLACEHOLDER)

    return {
        "injected": injected,
        "draft_path": str(draft.resolve()),
        "remaining_pending": load_state(state_path).get("pending", []),
    }


def _delete_paragraph(para) -> None:
    el = para._element
    el.getparent().remove(el)


def load_table_formats() -> dict[str, dict[str, Any]]:
    path = PACKAGE_SKILL_DIR / "table-formats.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def find_table_by_header(doc, header_prefix: str):
    """Return first table whose first cell starts with header_prefix (case-insensitive)."""
    want = header_prefix.strip().lower()
    for table in doc.tables:
        if not table.rows:
            continue
        first = table.rows[0].cells[0].text.strip().lower()
        if first.startswith(want):
            return table
    return None


def find_marked_table(doc, table_id: str):
    """Locate a table by ``table-formats.json`` marker, placeholder, or header row."""
    spec = load_table_formats().get(table_id)
    if not spec:
        return None
    cell_marker = str(spec.get("cell_marker", ""))
    if cell_marker:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell_marker in cell.text:
                        return table
    header_match = str(spec.get("header_match", ""))
    if header_match:
        found = find_table_by_header(doc, header_match)
        if found is not None:
            return found
    return None


def find_access_table(doc):
    return find_marked_table(doc, "access")


def find_revision_table(doc):
    return find_marked_table(doc, "revision")


def replace_marked_table(path: Path, table_id: str, rows: list[dict[str, str]]) -> dict[str, Any]:
    """Replace data rows in a marked table (header row preserved)."""
    spec = load_table_formats().get(table_id)
    if not spec:
        raise ValueError(f"Unknown table id: {table_id}")
    fields = list(spec["row_fields"])

    doc = open_document(path)
    table = find_marked_table(doc, table_id)
    if table is None:
        raise ValueError(
            f"Table '{table_id}' not found (expected header row starting with "
            f"'{spec.get('header_match')}'). Stage a pipe table first, then replace rows."
        )

    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)  # noqa: SLF001

    for row in rows:
        cells = table.add_row().cells
        for i, field in enumerate(fields):
            if i < len(cells):
                val = str(row.get(field, ""))
                para = cells[i].paragraphs[0]
                _fill_paragraph_with_inline_markdown(para, val, apply_body_font=True)

    doc.save(str(path))
    return {"table_id": table_id, "rows_written": len(rows), "path": str(path.resolve())}


def replace_access_table(path: Path, rows: list[dict[str, str]]) -> dict[str, Any]:
    """Replace rows in the ``access`` marked table (passwords / credentials)."""
    return replace_marked_table(path, "access", rows)


def append_marked_table_row(path: Path, table_id: str, row: dict[str, str]) -> dict[str, Any]:
    spec = load_table_formats().get(table_id)
    if not spec:
        raise ValueError(f"Unknown table id: {table_id}")
    doc = open_document(path)
    table = find_marked_table(doc, table_id)
    if table is None:
        raise ValueError(f"Table '{table_id}' not found.")
    cells = table.add_row().cells
    for i, field in enumerate(spec["row_fields"]):
        if i < len(cells):
            para = cells[i].paragraphs[0]
            _fill_paragraph_with_inline_markdown(para, str(row.get(field, "")), apply_body_font=True)
    doc.save(str(path))
    return {"table_id": table_id, "path": str(path.resolve()), **row}


def append_access_row(
    path: Path,
    *,
    system: str,
    purpose: str,
    username: str = "",
    password: str = "",
    notes: str = "",
) -> dict[str, Any]:
    return append_marked_table_row(
        path,
        "access",
        {
            "system": system,
            "purpose": purpose,
            "username": username,
            "password": password,
            "notes": notes,
        },
    )


def append_revision(
    path: Path,
    rev: str,
    date: str,
    description: str,
    author: str,
) -> dict[str, Any]:
    doc = open_document(path)
    table = find_revision_table(doc)
    if table is None:
        raise ValueError(
            "Revision table not found (expected first column header starting with 'Rev')."
        )
    row = table.add_row()
    cells = row.cells
    values = [rev, date, description, author]
    for i, val in enumerate(values):
        if i < len(cells):
            cells[i].text = val

    doc.save(str(path))
    return {"rev": rev, "date": date, "description": description, "author": author}


PACKAGE_SKILL_DIR = Path(__file__).resolve().parent.parent / "skills" / "template-docx-writer"

WORKSPACE_SEED_FILES = (
    ("section-catalog.md", "templates/section-catalog.md"),
    ("manual-styles.json", "templates/manual-styles.json"),
    ("table-formats.json", "templates/table-formats.json"),
    ("access-table-example-rows.json", "templates/access-table-example-rows.json"),
)


def ensure_manual_workspace(project_root: Path) -> dict[str, Any]:
    """Create workspace folders and seed templates/ from package defaults if missing."""
    root = project_root.resolve()
    templates = root / "templates"
    projects = root / "projects"
    templates.mkdir(parents=True, exist_ok=True)
    projects.mkdir(parents=True, exist_ok=True)

    seeded: list[str] = []
    for src_name, dest_rel in WORKSPACE_SEED_FILES:
        dest = root / dest_rel
        if dest.is_file():
            continue
        src = PACKAGE_SKILL_DIR / src_name
        if not src.is_file():
            continue
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dest)
        seeded.append(str(dest.resolve()))

    styles_path = (templates / "manual-styles.json").resolve()
    return {
        "workspace_root": str(root),
        "section_catalog": str((templates / "section-catalog.md").resolve()),
        "manual_styles": str(styles_path),
        "manual_styles_seeded_this_call": str(styles_path) in seeded,
        "package_default_styles": str(package_manual_styles_path().resolve()),
        "seeded_files": seeded,
    }


def start_project(
    project_root: Path,
    project_id: str,
    template_path: Path,
    *,
    fix_strict: bool = True,
) -> dict[str, Any]:
    template_path = template_path.resolve()

    # Template bundle: either a folder holding manifest.json + a .docx, or a .docx with
    # a sibling manifest.json. The bundle is the single source of truth for styles and
    # the section catalog (recorded as template_bundle_path in state.json). A bare
    # .docx with no sibling manifest falls back to legacy workspace templates/ seeding.
    bundle_dir: Path | None = None
    if template_path.is_dir():
        bundle_dir = template_path
    elif template_path.is_file() and (template_path.parent / "manifest.json").is_file():
        bundle_dir = template_path.parent

    if bundle_dir is not None:
        preferred = bundle_dir / "controls-template.docx"
        docx_path = preferred if preferred.is_file() else next(iter(sorted(bundle_dir.glob("*.docx"))), None)
        if docx_path is None:
            raise ValueError(f"No .docx in bundle folder: {bundle_dir}")
    elif template_path.is_file():
        docx_path = template_path
    else:
        raise ValueError(f"Template not found: {template_path}")

    workspace = ensure_manual_workspace(project_root)

    proj_dir = project_root / "projects" / project_id
    inputs = proj_dir / "inputs"
    output = proj_dir / "output"
    draft = proj_dir / "draft.docx"
    proj_dir.mkdir(parents=True, exist_ok=True)
    inputs.mkdir(exist_ok=True)
    output.mkdir(exist_ok=True)

    shutil.copy2(docx_path, draft)
    if fix_strict:
        import subprocess
        import sys

        fix_script = Path(__file__).resolve().parent / "fix_strict_ooxml.py"
        subprocess.run(
            [sys.executable, str(fix_script), str(draft), "--in-place"],
            check=True,
            capture_output=True,
            text=True,
        )

    placeholders = discover_placeholders(draft)
    state = {
        "project_id": project_id,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "template_path": str(docx_path),
        "draft_path": str(draft.resolve()),
        "placeholders": placeholders,
        "completed": [],
        "pending": placeholders.copy(),
        "access_table": "pending",
        "staging": {},
        "sources": [],
        "images": [],
    }
    if bundle_dir is not None:
        state["template_bundle_path"] = str(bundle_dir.resolve())
    state_path = proj_dir / "state.json"
    state_path.write_text(json.dumps(state, indent=2), encoding="utf-8")

    return {
        "project_id": project_id,
        "project_dir": str(proj_dir.resolve()),
        "draft_path": str(draft.resolve()),
        "state_path": str(state_path.resolve()),
        "placeholder_count": len(placeholders),
        "placeholders": placeholders,
        "workspace": workspace,
    }


def load_state(state_path: Path) -> dict[str, Any]:
    return json.loads(state_path.read_text(encoding="utf-8"))


def mark_placeholder_done(state_path: Path, placeholder: str) -> None:
    state = load_state(state_path)
    ph = placeholder.upper()
    if ph in state.get("pending", []):
        state["pending"] = [p for p in state["pending"] if p != ph]
    if ph not in state.get("completed", []):
        state["completed"] = [*state.get("completed", []), ph]
    state_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
